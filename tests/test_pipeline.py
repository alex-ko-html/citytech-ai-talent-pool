"""End-to-end monthly runs against offline fakes (tests/fakes.py)."""
import csv
import json
import shutil
from pathlib import Path

import openpyxl
import pytest

import fakes
import harness
from talentpool import inbox


def by_title(payload):
    return {i["title"]: i for i in payload["items"]}


def test_faculty_come_from_the_directory(tmp_path):
    p, _ = harness.run(tmp_path)
    names = {f["name"]: f for f in p["faculty"]}
    assert set(names) == {"Benito Mendoza", "Sunghoon Jang", "Viviana Acquaviva", "Lin Zhou"}   # no adjunct, no staff
    assert names["Benito Mendoza"]["school"] == "School of Technology & Design"
    cache = (tmp_path / "data" / "faculty_directory.csv").read_text()
    assert "718-000-0000" not in cache and "917-000-0000" not in cache                          # phones never stored


def test_types_sources_and_confidence(tmp_path):
    p, _ = harness.run(tmp_path)
    t = by_title(p)
    assert t["Large language models for data analysis"]["type"] == "Conference paper"
    assert t["Machine learning for ocean carbon"]["type"] == "Grant"
    assert t["Smart sensors for elders"]["amount"] == 210000                                   # NIH years summed
    advising = t["Towards AI-Driven Academic Advising: Lessons from a Domain-Specific Chatbot"]
    assert advising["src"] == ["OpenAlex", "Dept bio"]                                          # bio entry merged by DOI
    assert t["Teaching robotics with ROS 2"]["type"] == "Conference presentation"
    assert not t["Deep learning for flood forecasting in Florida"]["ok"]                        # same-name profile
    assert "Erratum to: An IoT-enabled environment" not in t                                    # excluded
    nsf_people = {p["faculty"][i["p"]]["name"] for i in p["items"] if i["title"] == "Machine learning for ocean carbon"}
    assert nsf_people == {"Viviana Acquaviva", "Benito Mendoza"}                                # PI + co-PI by email


def test_item_ids_are_permanent(tmp_path):
    p1, _ = harness.run(tmp_path)
    ct, other = fakes.default_works()
    ct.insert(0, fakes.work(50, "A brand new paper on chemistry", 2026, [fakes.author("Lin Zhou", "A5")]))
    p2, _ = harness.run(tmp_path, ct=ct, other=other)
    ids1 = {i["title"]: i["id"] for i in p1["items"]}
    ids2 = {i["title"]: i["id"] for i in p2["items"]}
    assert all(ids2[t] == i for t, i in ids1.items())
    assert ids2["A brand new paper on chemistry"] not in ids1.values()


def _edit_and_upload(tmp_path, edits, new_rows=()):
    src = tmp_path / "site" / "data" / "CityTech_Talent_Pool.xlsx"
    wb = openpyxl.load_workbook(src)
    for sheet in ("Items", "Items To Check"):
        ws = wb[sheet]
        hdr = {c.value: c.column for c in ws[1]}
        for row in ws.iter_rows(min_row=2):
            iid = row[hdr["ItemID"] - 1].value
            for (target, col), val in edits.items():
                if target == iid:
                    row[hdr[col] - 1].value = val
        if sheet == "Items":
            for nr in new_rows:
                ws.append([nr.get(c.value) for c in ws[1]])
    (tmp_path / "inbox").mkdir(exist_ok=True)
    wb.save(tmp_path / "inbox" / "reviewed.xlsx")
    return inbox.main(inbox=tmp_path / "inbox", data_dir=tmp_path / "data", overrides_path=tmp_path / "overrides.csv",
                      extra_path=tmp_path / "faculty_extra.csv", workbook=src)


def test_upload_round_trip(tmp_path):
    p, _ = harness.run(tmp_path)
    ids = {i["title"]: i["id"] for i in p["items"]}
    flood, yolo = ids["Deep learning for flood forecasting in Florida"], ids["YOLO bridge inspection with computer vision"]
    talk = ids["Teaching robotics with ROS 2"]
    summary = _edit_and_upload(tmp_path, {
        (flood, "Check"): "Not theirs", (yolo, "Check"): "Confirmed", (yolo, "Reviewed by"): "Pat Reviewer",
        (talk, "ItemType"): "Conference paper",
    }, new_rows=[{"Email": "SJang@citytech.cuny.edu", "FullName": "Sunghoon Jang", "ItemType": "Patent",
                  "ItemTitle": "Smart cane with fall detection", "Year": 2023, "Sources": "Manual"}])
    assert "3 review decisions, 1 items added" in summary[0]
    assert (tmp_path / "inbox" / "processed").exists() and not (tmp_path / "inbox" / "reviewed.xlsx").exists()
    p2, _ = harness.run(tmp_path)
    t = by_title(p2)
    assert "Deep learning for flood forecasting in Florida" not in t
    assert t["YOLO bridge inspection with computer vision"]["ok"] and t["YOLO bridge inspection with computer vision"]["rev"]
    assert t["Teaching robotics with ROS 2"]["type"] == "Conference paper"
    assert t["Smart cane with fall detection"]["type"] == "Patent"
    rows = list(csv.DictReader(open(tmp_path / "overrides.csv")))
    assert any(r["reviewed_by"] == "Pat Reviewer" and r["action"] == "confirm" for r in rows)   # who decided is kept


def test_import_outside_template_workbook(tmp_path):
    """A department's own workbook in the CUNY template (e.g. Google Scholar collection) merges in."""
    harness.run(tmp_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"
    from talentpool.workbook import ITEM_COLUMNS
    ws.append(ITEM_COLUMNS)
    row = dict(ItemID="CET-0001", LastName="Mendoza", FirstName="Benito", FullName="Benito Mendoza",
               Email="BMendoza@citytech.cuny.edu", ItemType="Journal article", Year=2025, Sources="Dept bio; Google Scholar",
               ItemTitle="Towards AI-Driven Academic Advising: Lessons from a Domain-Specific", ScholarCitations=3,
               Confidence="High")
    ws.append([row.get(c) for c in ITEM_COLUMNS])
    row2 = dict(row, ItemID="CET-0090", ItemTitle="A Scholar-only paper about IoT gateways", Sources="Google Scholar",
                Confidence="Medium (Google Scholar profile)")
    ws.append([row2.get(c) for c in ITEM_COLUMNS])
    f = wb.create_sheet("Faculty")
    f.append(["FullName", "Email", "ORCID", "ScholarProfile", "Hindex"])
    f.append(["Benito Mendoza", "BMendoza@citytech.cuny.edu", "0000-0002-1825-0097", "https://scholar.google.com/x", 9])
    (tmp_path / "inbox").mkdir()
    wb.save(tmp_path / "inbox" / "cet_pilot.xlsx")
    summary = inbox.main(inbox=tmp_path / "inbox", data_dir=tmp_path / "data", overrides_path=tmp_path / "overrides.csv",
                         extra_path=tmp_path / "faculty_extra.csv", workbook=tmp_path / "site/data/CityTech_Talent_Pool.xlsx")
    assert "template workbook" in summary[0] and "2 items added" in summary[0]                  # its IDs are NOT treated as ours
    p, _ = harness.run(tmp_path)
    t = by_title(p)
    assert "Google Scholar" in t["Towards AI-Driven Academic Advising: Lessons from a Domain-Specific Chatbot"]["src"]
    assert t["A Scholar-only paper about IoT gateways"]["src"] == ["Google Scholar"]
    fac = openpyxl.load_workbook(tmp_path / "site/data/CityTech_Talent_Pool.xlsx")["Faculty"]
    hdr = [c.value for c in fac[1]]
    m = [dict(zip(hdr, [c.value for c in r])) for r in fac.iter_rows(min_row=2) if r[3].value == "Benito Mendoza"][0]
    assert m["ORCID"] == "0000-0002-1825-0097" and m["Hindex"] == 9


def test_directory_outage_uses_last_copy(tmp_path):
    harness.run(tmp_path)
    p, rep = harness.run(tmp_path, directory_up=False)
    assert len(p["faculty"]) == 4
    assert any("last month's copy" in w for w in rep["warnings"])


def test_broken_run_is_refused(tmp_path):
    harness.run(tmp_path)
    good = (tmp_path / "site/data/talent_pool.json").read_text()
    with pytest.raises(SystemExit):
        harness.run(tmp_path, ct=[], other=[])
    assert (tmp_path / "site/data/talent_pool.json").read_text() == good


# Column order of CUNY_Talent_Pool_Publication_Data_Structure.xlsx (the template, Oct 2026).
TEMPLATE = {
    "Items": ["ItemID", "LastName", "FirstName", "FullName", "Title", "College", "Department", "Email", "ItemType",
              "ItemTitle", "Year", "Venue / Funder", "Authors / Investigators", "Role", "Amount (USD)",
              "AwardID / PatentNo", "DOI", "URL", "Sources", "SourceCount", "Confidence", "ScholarCitations",
              "CitationAsListed", "Abstract", "AbstractSource", "FullTextURL", "FullTextType", "OpenAccessStatus"],
    "Faculty": ["SNo", "LastName", "FirstName", "FullName", "Title", "College", "Department", "Email", "Office",
                "DirectoryProfile", "ORCID", "HasProfile", "TotalPubs", "Citations", "Hindex", "Iindex",
                "ScholarProfile", "ScholarInterests", "ScholarCheckedOn", "ItemsCollected", "Journal article",
                "Conference paper", "Book", "Book chapter", "Conference presentation", "Technical report",
                "Paper (venue type unclear)", "Patent", "Grant", "WithDOI", "FoundIn: Dept bio",
                "FoundIn: Academic Works", "FoundIn: ORCID", "FoundIn: Crossref", "FoundIn: NSF",
                "FoundIn: Google Scholar", "Notes"],
    "Sources": ["Source", "Endpoint", "What it gave", "Status"],
}


def test_workbook_keeps_the_template_columns(tmp_path):
    harness.run(tmp_path)
    wb = openpyxl.load_workbook(tmp_path / "site/data/CityTech_Talent_Pool.xlsx")
    assert wb.sheetnames[:3] == ["Items", "Faculty", "Sources"] and "Items Excluded" in wb.sheetnames
    for sheet, want in TEMPLATE.items():
        assert [c.value for c in wb[sheet][1]][:len(want)] == want, sheet


def test_items_to_check_stay_manageable(tmp_path):
    ct, other = fakes.default_works()
    other.append(fakes.work(40, "Soil chemistry of rice paddies", 2019, [fakes.author("Lin Zhou", "A5", False)]))
    p, _ = harness.run(tmp_path, ct=ct, other=other)
    titles = {i["title"] for i in p["items"]}
    assert "Soil chemistry of rice paddies" not in titles                    # not AI-related: not worth checking
    assert "YOLO bridge inspection with computer vision" in titles           # AI-related: to check
    many = [fakes.work(1000 + k, f"Deep learning study {k}", 2020, [fakes.author("Lin Zhou", "A5", False)]) for k in range(160)]
    p, rep = harness.run(tmp_path, ct=ct, other=other + many)
    lin = [i for i in p["items"] if p["faculty"][i["p"]]["name"] == "Lin Zhou"]
    assert all(i["ok"] for i in lin)                                         # merged profile set aside entirely
    assert any("probably several people" in w for w in rep["warnings"])


def test_openalex_types():
    from talentpool.items import type_from_openalex as t
    assert t("conference-paper", "conference", "ASEE Annual Conference and Exposition") == "Conference paper"
    assert t("conference-abstract", "", "ECS Meeting Abstracts") == "Conference presentation"
    assert t("some-new-type", "journal", "Journal X") == "Journal article"        # unknown types use the venue
    assert t("other", "", "IEEE Frontiers in Education Conference (FIE)") == "Conference paper"
    assert t("erratum") is None
