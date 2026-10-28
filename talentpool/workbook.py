"""Writes the Excel deliverable in the CUNY Talent Pool template format, and reads it back
when someone uploads an edited copy (see inbox.py).

Tabs and columns follow CUNY_Talent_Pool_Publication_Data_Structure.xlsx exactly; columns this
project adds (School, AI flags, review columns) come after the template's own columns.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .items import ITEM_TYPES, EXTRA_TYPES

ITEM_COLUMNS = ["ItemID", "LastName", "FirstName", "FullName", "Title", "College", "Department", "Email",
                "ItemType", "ItemTitle", "Year", "Venue / Funder", "Authors / Investigators", "Role",
                "Amount (USD)", "AwardID / PatentNo", "DOI", "URL", "Sources", "SourceCount", "Confidence",
                "ScholarCitations", "CitationAsListed", "Abstract", "AbstractSource", "FullTextURL",
                "FullTextType", "OpenAccessStatus"]
ITEM_EXTRA = ["School", "Citations (OpenAlex)", "AI-related", "AI evidence", "Check", "Correction notes",
              "Reviewed by"]
FACULTY_COLUMNS = ["SNo", "LastName", "FirstName", "FullName", "Title", "College", "Department", "Email",
                   "Office", "DirectoryProfile", "ORCID", "HasProfile", "TotalPubs", "Citations", "Hindex",
                   "Iindex", "ScholarProfile", "ScholarInterests", "ScholarCheckedOn", "ItemsCollected",
                   *ITEM_TYPES, "WithDOI", "FoundIn: Dept bio", "FoundIn: Academic Works", "FoundIn: ORCID",
                   "FoundIn: Crossref", "FoundIn: NSF", "FoundIn: Google Scholar", "Notes"]
FACULTY_EXTRA = ["School", "Status", *EXTRA_TYPES, "FoundIn: OpenAlex", "FoundIn: NIH", "AI tier",
                 "AI items (confirmed)", "Items to check"]
SOURCE_COLUMNS = ["Source", "Endpoint", "What it gave", "Status"]
CHECK_CHOICES = ["Confirmed", "Not theirs", "Not AI", "Is AI"]
EDITABLE = {"ItemType": "item_type", "ItemTitle": "title", "Year": "year", "Venue / Funder": "venue",
            "Role": "role", "Amount (USD)": "amount"}

FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
EXTRA_FILL = PatternFill("solid", fgColor="44617F")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")       # cells reviewers fill in
WIDTHS = {"ItemTitle": 60, "Abstract": 50, "CitationAsListed": 50, "Authors / Investigators": 36,
          "Venue / Funder": 32, "Department": 28, "Confidence": 34, "URL": 30, "FullTextURL": 30,
          "AI evidence": 36, "Correction notes": 30, "ScholarInterests": 30, "Notes": 30, "Sources": 26}


def _item_row(it, p):
    return [it["item_id"], p["last_name"], p["first_name"], p["full_name"], p["title"], p["college"],
            p["department"], p["email"], it["item_type"], it["title"], it["year"] or None, it["venue"],
            it["authors"], it["role"], _num(it["amount"]), it["award_id"], it["doi"], it["url"],
            "; ".join(it["sources"]), len(it["sources"]), it["confidence"], _num(it["scholar_citations"]),
            it["citation_as_listed"], it["abstract"][:32000], it["abstract_source"], it["fulltext_url"],
            it["fulltext_type"], it["oa_status"],
            p["school"], _num(it["citations"]), "Yes" if it.get("is_ai") else "No", it.get("ai_evidence", ""),
            {"confirmed": "Confirmed"}.get(it.get("review"), ""), "", it.get("reviewed_by", "")]


def _num(v):
    if v in ("", None):
        return None
    try:
        f = float(str(v).replace(",", "").replace("$", ""))
        return int(f) if f.is_integer() else f
    except ValueError:
        return v


def _sheet(wb, title, cols, extra=(), first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for i, c in enumerate([*cols, *extra], 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HEAD, (HEAD_FILL if i <= len(cols) else EXTRA_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 14)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    return ws


def _finish(ws, nrows):
    if nrows:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{nrows + 1}"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT


def _review_columns(ws, nrows):
    col = {c.value: c.column_letter for c in ws[1]}
    dv = DataValidation(type="list", formula1='"' + ",".join(CHECK_CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for name in ("Check", "Correction notes", "Reviewed by"):
        ws[f"{col[name]}1"].comment = Comment(
            {"Check": "Pick: Confirmed (this item is theirs and correct), Not theirs (belongs to someone else), "
                      "Not AI / Is AI (fix the AI flag). Then upload the file.",
             "Correction notes": "Anything else to fix. You can also edit ItemType, ItemTitle, Year, "
                                 "Venue / Funder, Role or Amount directly in their cells.",
             "Reviewed by": "Your name, so the next maintainer knows who decided."}[name], "Talent pool")
        for r in range(2, nrows + 2):
            ws[f"{col[name]}{r}"].fill = INPUT_FILL
    if nrows:
        dv.add(f"{col['Check']}2:{col['Check']}{nrows + 1}")


def write_workbook(path, people, items, tiers, source_rows, generated):
    by_key = {p["key"]: p for p in people}
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    groups = {"Items": [], "Items To Check": [], "Items Excluded": []}
    for it in sorted(items, key=lambda i: (i["item_id"])):
        p = by_key[it["person"]]
        if it["excluded_reason"]:
            groups["Items Excluded"].append(_item_row(it, p) + [it["excluded_reason"]])
        elif it["confidence"].startswith("Low"):
            groups["Items To Check"].append(_item_row(it, p))
        else:
            groups["Items"].append(_item_row(it, p))

    ws = _sheet(wb, "Items", ITEM_COLUMNS, ITEM_EXTRA, first=True)
    for r in groups["Items"]:
        ws.append(r)
    _finish(ws, len(groups["Items"]))
    _review_columns(ws, len(groups["Items"]))

    # Faculty: counts are formulas over the Items tab, so they follow a reviewer's edits.
    wf = _sheet(wb, "Faculty", FACULTY_COLUMNS, FACULTY_EXTRA)
    col = {c: get_column_letter(i) for i, c in enumerate(ITEM_COLUMNS + ITEM_EXTRA, 1)}
    n_items = len(groups["Items"]) + 1
    rng = lambda c: f"Items!${col[c]}$2:${col[c]}${max(n_items, 2)}"
    faculty = sorted([p for p in people if p.get("_faculty")], key=lambda p: (p["school"], p["department"], p["last_name"]))
    to_check = {}
    for it in items:
        if it["confidence"].startswith("Low") and not it["excluded_reason"]:
            to_check[it["person"]] = to_check.get(it["person"], 0) + 1
    for n, p in enumerate(faculty, 1):
        r = n + 1
        who = f'{rng("Email")},$H{r}' if p["email"] else f'{rng("FullName")},$D{r}'
        cnt = lambda cond="": f"=COUNTIFS({who}{cond})"
        t = tiers.get(p["key"], {})
        row = [n, p["last_name"], p["first_name"], p["full_name"], p["title"], p["college"], p["department"],
               p["email"], p["office"], p["directory_profile"], p["orcid"], bool(p.get("scholar_profile")),
               _num(p.get("scholar_total_pubs")), _num(p.get("scholar_citations")), _num(p.get("hindex")),
               _num(p.get("iindex")), p.get("scholar_profile", ""), p.get("scholar_interests", ""),
               p.get("scholar_checked_on", ""), cnt()]
        row += [cnt(f',{rng("ItemType")},"{ty}"') for ty in ITEM_TYPES]
        row += [cnt(f',{rng("DOI")},"<>"')]
        for s in ("Dept bio", "CUNY Academic Works", "ORCID", "Crossref", "NSF", "Google Scholar"):
            row.append(cnt(f',{rng("Sources")},"*{s}*"'))
        row += [p.get("notes", ""), p["school"], p["status"]]
        row += [cnt(f',{rng("ItemType")},"{ty}"') for ty in EXTRA_TYPES]
        row += [cnt(f',{rng("Sources")},"*OpenAlex*"'), cnt(f',{rng("Sources")},"*NIH*"'),
                t.get("tier", ""), cnt(f',{rng("AI-related")},"Yes"'), to_check.get(p["key"], 0)]
        wf.append(row)
    _finish(wf, len(faculty))

    wsrc = _sheet(wb, "Sources", SOURCE_COLUMNS)
    for r in source_rows:
        wsrc.append(list(r))
    wsrc.column_dimensions["B"].width, wsrc.column_dimensions["C"].width = 60, 60
    _finish(wsrc, len(source_rows))

    for name, extra in (("Items Excluded", ["Reason excluded"]), ("Items To Check", [])):
        w = _sheet(wb, name, ITEM_COLUMNS, ITEM_EXTRA + extra)
        for r in groups[name]:
            w.append(r)
        _finish(w, len(groups[name]))
        if name == "Items To Check":
            _review_columns(w, len(groups[name]))

    how = wb.create_sheet("How to use")
    how.column_dimensions["A"].width = 110
    lines = [
        ("CUNY Talent Pool: City Tech publication data", True),
        (f"Generated {generated} by the monthly update. Do not reorder or rename columns.", False),
        ("", False),
        ("Tabs", True),
        ("Items: one row per publication, presentation, grant or patent, with High/Medium confidence.", False),
        ("Faculty: one row per full-time faculty member. Counts are formulas over the Items tab.", False),
        ("Sources: which databases this run used and what each returned.", False),
        ("Items Excluded: records that aren't research items, or that a reviewer marked 'Not theirs'.", False),
        ("Items To Check: same-name matches with no City Tech affiliation. Confirm or reject them.", False),
        ("", False),
        ("To review or correct (yellow cells)", True),
        ("1. In the Check column pick Confirmed, Not theirs, Not AI or Is AI.", False),
        ("2. Fix ItemType, ItemTitle, Year, Venue / Funder, Role or Amount directly in their cells if wrong.", False),
        ("3. To add a missing item, add a row at the bottom of Items with an empty ItemID (fill name, "
         "Email, ItemType, ItemTitle, Year, Sources).", False),
        ("4. Put your name in Reviewed by, save, and upload the file with the dashboard's Upload button.", False),
        ("The next update applies your changes and keeps them every month.", False),
        ("", False),
        ("Confidence", True),
        ("High: the paper lists City Tech, or it's in CUNY Academic Works, an NSF/NIH award to City Tech, "
         "their ORCID record, their faculty bio, or a reviewer confirmed it.", False),
        ("Medium: found in two sources, or a Google Scholar profile that lists City Tech.", False),
        ("Low: only on an author profile with the same name. Often someone else.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = how.cell(row=i, column=1, value=text)
        c.font = Font(name="Arial", size=12 if i == 1 else 10, bold=bold)
    wb.save(path)


def read_workbook(path):
    """Returns {"Items": [dict...], "Items To Check": [...], "Faculty": [...]} from an uploaded copy."""
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for name in ("Items", "Items To Check", "Faculty"):
        if name not in wb.sheetnames:
            continue
        rows = wb[name].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        out[name] = [{h: ("" if v is None else v) for h, v in zip(header, r) if h}
                     for r in rows if any(v not in (None, "") for v in r)]
    return out
